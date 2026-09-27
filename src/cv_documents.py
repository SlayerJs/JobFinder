"""Render the same structured CV to DOCX and paginated PDF."""
import os
import uuid
from pathlib import Path
from xml.sax.saxutils import escape

from src.cv import one


def render_version(service, version_id, extension):
    if extension not in {'pdf', 'docx', 'tex'}:
        raise ValueError('Use PDF, DOCX or TEX')
    version = one(service.db, 'versions', version_id)
    category = one(service.db, 'categories', version['category_id'])
    profile = one(service.db, 'profiles', category['profile_id'])
    source = one(service.db, 'sources', profile['source_id'])
    data = version['data']
    if extension == 'tex' and not data.get('latex_source'):
        raise ValueError('This revision has no LaTeX template. Import a .tex CV to generate LaTeX versions.')
    path = service.private_dir / f"cv-category-{category['id']}-r{version_id}.{extension}"
    if path.exists():
        return path
    temporary_path = path.with_name(uuid.uuid4().hex + '.' + extension)
    if extension == 'tex':
        temporary_path.write_bytes(data['latex_source'].encode('utf-8'))
    elif extension == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt
        document = Document()
        document.sections[0].top_margin = Inches(.65)
        document.sections[0].bottom_margin = Inches(.65)
        document.styles['Normal'].font.name = 'DejaVu Sans'
        document.styles['Normal'].font.size = Pt(10)
        document.add_heading(data['title'], 0)
        if source['contacts']:
            document.add_paragraph(source['contacts'])
        for section in data['sections']:
            document.add_heading(section['heading'], 1)
            for item in section['items']:
                document.add_paragraph(item['text'])
        document.save(temporary_path)
    else:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        font = os.getenv('JOBFINDER_CV_FONT')
        candidates = [Path(font)] if font else [
            Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
            Path('/usr/share/fonts/TTF/DejaVuSans.ttf'),
            Path('/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf'),
            Path('C:/Windows/Fonts/arial.ttf'),
            Path('/mnt/c/Windows/Fonts/arial.ttf'),
            Path('/Library/Fonts/Arial Unicode.ttf')]
        font_path = next((p for p in candidates if p.is_file()), None)
        if font_path is None:
            raise ValueError('Set JOBFINDER_CV_FONT to a Unicode TrueType font for PDF exports')
        font_name = 'CVUnicode-' + str(abs(hash(str(font_path))))
        if font_name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
        face = pdfmetrics.getFont(font_name).face
        text = data['title'] + source['contacts'] + ''.join(
            s['heading'] + ''.join(i['text'] for i in s['items']) for s in data['sections'])
        if any(ord(char) not in face.charToGlyph for char in text if not char.isspace()):
            raise ValueError('The selected PDF font lacks some characters; set JOBFINDER_CV_FONT to a font covering this CV language')
        normal = ParagraphStyle('CV', fontName=font_name, fontSize=10, leading=14,
                                spaceAfter=7, alignment=TA_LEFT, splitLongWords=True)
        heading = ParagraphStyle('CVHeading', parent=normal, fontSize=13, leading=17,
                                 spaceBefore=12, spaceAfter=6, keepWithNext=True, textColor=colors.HexColor('#19334d'))
        title = ParagraphStyle('CVTitle', parent=heading, fontSize=19, leading=23)
        paragraph = lambda value, style: Paragraph(escape(value).replace('\n', '<br/>'), style)
        story = [paragraph(data['title'], title)]
        if source['contacts']:
            story.extend([paragraph(source['contacts'], normal), Spacer(1, 8)])
        for section in data['sections']:
            story.append(paragraph(section['heading'], heading))
            story.extend(paragraph(item['text'], normal) for item in section['items'])
        def footer(canvas, document):
            canvas.setFont(font_name, 8)
            canvas.drawRightString(A4[0]-42, 24, str(document.page))
        SimpleDocTemplate(str(temporary_path), pagesize=A4, leftMargin=42, rightMargin=42,
                          topMargin=35, bottomMargin=40).build(story, onFirstPage=footer, onLaterPages=footer)
    temporary_path.chmod(0o600)
    temporary_path.replace(path)
    return path
