import contextlib
import io
import os
import tempfile
import unittest
from unittest.mock import Mock, patch

import main
from src.ai_filter import AIFilter
from src.database import JobDatabase
from src.models import JobPosting
from src.configuration import load_config
from pathlib import Path


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.config_patch = patch.object(main, "USER_CONFIG", load_config(Path(__file__).resolve().parents[1] / "config.example.yaml"))
        self.config_patch.start()
        self.addCleanup(self.config_patch.stop)
        self.db = JobDatabase(':memory:')
        self.jobs = [JobPosting(str(i),f'Stage {i}','Example','France',f'https://example.com/{i}',
                     (f'Stage cloud security project {i} in France. '*15),'Test') for i in range(3)]
        for job in self.jobs:
            self.db.add_job(job)

    def tearDown(self):
        self.db.close()

    def test_process_persists_decisions_and_report(self):
        ai = AIFilter({'batch_input_tokens':20000})
        ai.evaluate_batch = Mock(return_value={
            '0':{'decision':'APPROVED','tier':4,'reason_code':'MATCH','reason':'Suitable','evidence':'Stage cloud'},
            '1':{'decision':'REJECTED','tier':None,'reason_code':'CONTRACT','reason':'Wrong contract','evidence':'project 1'}})
        notifier = Mock()
        notifier.send_job_alert.return_value = True
        old_cwd = os.getcwd()
        with tempfile.TemporaryDirectory() as directory:
            try:
                os.chdir(directory)
                with patch.multiple(main,db=self.db,ai=ai,notifier=notifier):
                    main.process_pending_jobs()
                states = {r['id']:r['status'] for r in self.db.conn.execute('SELECT id,status FROM jobs')}
                self.assertEqual(states,{'0':'APPROVED','1':'REJECTED','2':'PENDING'})
                notifier.send_job_alert.assert_called_once()
                notifier.send_run_summary.assert_called_once()
                report = notifier.send_run_summary.call_args.args[0]
                from openpyxl import load_workbook
                workbook = load_workbook(report)
                self.assertEqual(workbook.active.max_row,4)
                workbook.close()
                self.assertEqual(len(self.db.pending_notifications()),0)
            finally:
                os.chdir(old_cwd)

    def test_dry_run_has_no_external_calls_or_status_changes(self):
        ai = AIFilter()
        ai.evaluate_batch = Mock(side_effect=AssertionError('API forbidden'))
        notifier = Mock()
        with patch.multiple(main,db=self.db,ai=ai,notifier=notifier), contextlib.redirect_stdout(io.StringIO()) as output:
            main.process_pending_jobs(dry_run=True)
        self.assertIn('estimated_input_tokens',output.getvalue())
        self.assertEqual(len(self.db.get_pending_jobs()),3)
        ai.evaluate_batch.assert_not_called()
        self.assertFalse(notifier.mock_calls)
