import copy
import io
import json
import sqlite3
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace as NS
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas

from src.ai_filter import AIFilter
from src.cv import (CVService, active_version, categories, current_profile, encode, extract_cv,
                    matching_context, one, recommendation, rows, save_source, source_facts,
                    store_match, validate_analysis, validate_version)
from src.cv_documents import render_version
from src.cv_worker import CVWorker
from src.database import JobDatabase
from src.models import JobPosting


SOURCE = 'Ingénieure logiciel\nAcme — 2022–2024 — Développement Python\nMaster informatique — 2022'


class FakeProvider:
    """Deterministic response fixtures, no paid calls."""
    def __init__(self):
        self.calls = []
        self.client = NS(chat=NS(completions=NS(create=self.create)))

    def create(self, **kwargs):
        self.calls.append(kwargs)
        prompt = kwargs['messages'][0]['content']
        payload = json.loads(kwargs['messages'][1]['content'])
        if prompt.startswith('Analyze'):
            facts = payload['facts']
            data = {'profile': {'language': 'Français', 'summary': 'Expérience en développement Python',
                    'fact_ids': [f['id'] for f in facts], 'facts': facts},
                    'categories': [{'name': 'Développement logiciel', 'rationale': 'Expérience Python chez Acme',
                                    'fit': 'strong', 'role_titles': ['Développeuse Python'], 'gaps': [], 'fact_ids': ['f2']}]}
        elif prompt.startswith('Tailor'):
            data = {'language': payload['profile']['language'], 'title': 'Curriculum vitæ',
                    'sections': [{'heading': 'Expérience et formation', 'items': [
                        {'text': f['text'], 'fact_ids': [f['id']]} for f in payload['facts']]}],
                    'changes': ['Expérience mise en avant'], 'suggestions': ['Préciser les projets si documentés']}
            if 'Output format: standalone LaTeX' in prompt:
                data['latex_layout'] = {'font_size': 11, 'margin_mm': 18}
        elif prompt.startswith('Classify'):
            context = json.loads(prompt.split('Context (shared once for this batch):\n')[1])
            data = {'results': [{'id': j['id'], 'category_id': context['categories'][0]['id'],
                                 'matching_reason': 'Expérience Python pertinente'} for j in payload]}
        else:
            data = {'results': [{'id': j['id'], 'decision': 'APPROVED', 'tier': 2,
                    'reason_code': 'MATCH', 'reason': 'Suitable', 'evidence': 'Python',
                    'category_id': 99999, 'matching_reason': 'Invalid category fixture'} for j in payload]}
        return NS(usage=NS(prompt_tokens=100, completion_tokens=80, prompt_cache_hit_tokens=20),
                  choices=[NS(finish_reason='stop', message=NS(content=encode(data)))])


def docx_fixture(text=SOURCE, table=True):
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    if table:
        row = document.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = 'Compétences'
        row.cells[1].text = 'Python'
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class CVTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.db = JobDatabase(':memory:')
        self.addCleanup(self.db.close)
        self.provider = FakeProvider()
        self.ai = AIFilter({'max_retries': 0, 'batch_input_tokens': 30000}, self.db, self.provider.client)
        self.service = CVService(self.db, self.ai, self.temp.name)
        self.sid = save_source(self.db, 'cv.docx', SOURCE, 'Élise Example\nelise@example.test')

    def prepared(self):
        pid = self.service.analyze(self.sid)
        cats = categories(self.db)
        ids = self.service.save_categories(pid, [{**c['data'], 'id': c['id'], 'approved': True} for c in cats])
        vid = self.service.generate(ids[0])
        return pid, ids[0], vid

    def add_job(self, ident='job1', status='APPROVED'):
        job = JobPosting(ident, 'Python engineer', 'Example', 'Paris', 'https://example.test/' + ident,
                         (ident + ' Python cloud development in France. ') * 20, 'fixture')
        self.db.add_job(job)
        self.db.update_job_status(ident, status, 'Eligibility preserved', '2')
        self.db.notification_sent(ident)
        return job

    def test_extraction_docx_tables_accents_pdf_and_empty(self):
        text = extract_cv('cv.docx', docx_fixture())
        self.assertIn('Ingénieure', text)
        self.assertIn('Compétences | Python', text)
        self.assertLess(text.index('Master'), text.index('Compétences'))
        pdf = io.BytesIO()
        canvas = Canvas(pdf)
        canvas.drawString(30, 700, 'CV Python engineer')
        canvas.save()
        self.assertIn('Python engineer', extract_cv('cv.pdf', pdf.getvalue()))
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        blank = io.BytesIO()
        writer.write(blank)
        self.assertEqual(extract_cv('scanned.pdf', blank.getvalue()), '')
        self.assertEqual(extract_cv('empty.docx', docx_fixture('', False)), '')
        with self.assertRaises(ValueError):
            save_source(self.db, 'empty.pdf', '')
        for name, content in [('cv.txt', b'text'), ('cv.pdf', b'broken'), ('cv.docx', b'')]:
            with self.assertRaises(ValueError):
                extract_cv(name, content)

    def test_scanned_pdf_requires_replacement(self):
        from PIL import Image
        from reportlab.lib.utils import ImageReader
        image = Image.new('RGB', (200, 100), 'white')
        stream = io.BytesIO()
        canvas = Canvas(stream)
        canvas.drawImage(ImageReader(image), 20, 20)
        canvas.save()
        result = self.service.import_file('scan.pdf', stream.getvalue())
        self.assertEqual(result['text'], '')
        self.assertIn('OCR', result['notice'])

    def test_corrections_privacy_and_cache(self):
        sid = save_source(self.db, 'corrected.docx', SOURCE.replace('2024', '2025'), 'private@example.test')
        pid = self.service.analyze(sid)
        call = self.provider.calls[0]
        self.assertIn('2025', call['messages'][1]['content'])
        self.assertNotIn('private@example', encode(self.provider.calls))
        self.assertEqual(self.service.analyze(sid), pid)
        self.assertTrue(self.service.estimate('analyze', {'source_id': sid})['cached'])
        self.assertEqual(len(self.provider.calls), 1)
        profile = one(self.db, 'profiles', pid)['data']
        profile['summary'] = 'Résumé corrigé'
        edited = self.service.edit_profile(pid, profile)
        self.assertEqual(self.service.analyze(sid), edited)
        self.assertEqual(len(self.provider.calls), 1)

    def test_evidence_validation_and_unknown_categories(self):
        pid, cid, vid = self.prepared()
        source = one(self.db, 'sources', self.sid)
        profile = one(self.db, 'profiles', pid)['data']
        bad = copy.deepcopy(profile)
        bad['facts'][0]['text'] = 'Invented experience'
        with self.assertRaises(ValueError):
            validate_analysis({'profile': bad, 'categories': []}, source)
        bad = one(self.db, 'versions', vid)['data']
        bad['sections'][0]['items'][0]['fact_ids'] = ['unknown']
        with self.assertRaises(ValueError):
            validate_version(bad, source, 'Français')
        bad = one(self.db, 'versions', vid)['data']
        bad['sections'][0]['items'][0]['text'] = 'Saved 999 million euros'
        with self.assertRaises(ValueError):
            validate_version(bad, source, 'Français')
        bad['sections'][0]['items'][0]['text'] = 'Expert cloud architect'
        self.assertTrue(validate_version(bad, source, 'Français')['warnings'])
        self.add_job()
        with self.assertRaises(ValueError):
            store_match(self.db, 'job1', {'category_id': 9876, 'matching_reason': 'Unknown'})

    def test_approval_edit_history_and_rejected_job(self):
        _, cid, vid = self.prepared()
        self.add_job()
        store_match(self.db, 'job1', {'category_id': cid, 'matching_reason': 'Python'})
        self.assertEqual(recommendation(self.db, 'job1')['cv'], 'CV awaiting approval')
        self.assertIsNone(active_version(self.db, cid))
        with self.assertRaises(ValueError):
            self.service.approve(vid, False)
        self.service.approve(vid, True)
        self.assertEqual(recommendation(self.db, 'job1')['version_id'], vid)
        data = one(self.db, 'versions', vid)['data']
        data['title'] = 'CV révisé'
        new = self.service.edit_version(vid, data)
        self.assertNotEqual(new, vid)
        self.assertEqual(active_version(self.db, cid), vid)
        self.assertNotEqual(one(self.db, 'versions', vid)['data']['title'], data['title'])
        self.service.approve(new, True)
        self.assertEqual(active_version(self.db, cid), new)
        self.assertEqual(recommendation(self.db, 'job1')['version_id'], new)
        history = rows(self.db, 'SELECT * FROM cv_matches ORDER BY id')
        self.assertEqual(len(history), 3)
        self.assertEqual(history[1]['version_id'], vid)
        self.db.update_job_status('job1', 'REJECTED', 'Now rejected')
        self.assertIsNone(recommendation(self.db, 'job1')['version_id'])

    def test_source_and_category_staleness(self):
        pid, cid, vid = self.prepared()
        self.service.approve(vid, True)
        category = one(self.db, 'categories', cid)['data']
        category['rationale'] += ' edited'
        new_ids = self.service.save_categories(pid, [{**category, 'id': cid, 'approved': True}])
        self.assertNotEqual(cid, new_ids[0])
        self.assertTrue(one(self.db, 'versions', vid)['stale'])
        with self.assertRaises(ValueError):
            self.service.approve(vid, True)
        save_source(self.db, 'new.docx', SOURCE + '\nSQL')
        self.assertIsNone(current_profile(self.db))
        self.assertEqual(categories(self.db), [])

    def test_generation_cache_and_operations_accounting(self):
        _, cid, vid = self.prepared()
        self.assertEqual(self.service.generate(cid), vid)
        self.assertEqual(len(self.provider.calls), 2)
        usage = rows(self.db, 'SELECT * FROM api_usage')
        self.assertEqual([r['operation'] for r in usage], ['profile_analysis', 'cv_generation'])
        self.assertTrue(all(r['input_tokens'] == 100 and r['output_tokens'] == 80 for r in usage))
        self.assertEqual(self.service.estimate('generate', {'category_id': cid})['requests'], 0)

    def test_prompt_revision_invalidates_cache_on_explicit_operation(self):
        pid, cid, vid = self.prepared()
        with patch('src.cv.PROMPT_REVISION', 'cv-next'):
            self.assertFalse(self.service.estimate('analyze', {'source_id': self.sid})['cached'])
            new_pid = self.service.analyze(self.sid)
            self.assertNotEqual(pid, new_pid)
            self.assertTrue(one(self.db, 'versions', vid)['stale'])
            self.assertEqual(len(self.provider.calls), 3)

    def test_backfill_batches_cache_preserves_eligibility_and_alert_history(self):
        _, cid, vid = self.prepared()
        self.service.approve(vid, True)
        self.ai.config['max_jobs_per_batch'] = 2
        for i in range(5):
            self.add_job(str(i))
        self.add_job('rejected', 'REJECTED')
        before = rows(self.db, 'SELECT * FROM jobs')
        evaluations = rows(self.db, 'SELECT * FROM evaluations')
        self.assertEqual(self.service.classify()['matched'], 5)
        self.assertEqual(rows(self.db, 'SELECT * FROM jobs'), before)
        self.assertEqual(rows(self.db, 'SELECT * FROM evaluations'), evaluations)
        self.assertEqual(recommendation(self.db, '0')['version_id'], vid)
        calls = self.provider.calls[2:]
        self.assertEqual(len(calls), 3)
        for call in calls:
            prompt = call['messages'][0]['content']
            self.assertEqual(prompt.count('Context (shared once for this batch)'), 1)
            self.assertNotIn('Master informatique', prompt)
            self.assertNotIn('elise@example', encode(call))
            self.assertLessEqual(len(json.loads(call['messages'][1]['content'])), 2)
        self.assertEqual(self.service.classify()['matched'], 0)
        self.assertEqual(len(self.provider.calls), 5)

    def test_matching_failure_does_not_lose_eligibility(self):
        self.prepared()
        job = self.add_job()
        result = self.ai.evaluate_batch([job], 'Python roles')[job.id]
        self.assertEqual(result['decision'], 'APPROVED')
        self.assertEqual(result['category_id'], 99999)
        with self.assertRaises(ValueError):
            store_match(self.db, job.id, result)
        self.assertEqual(self.service.classify()['matched'], 1)

    def test_malformed_retry_and_budget(self):
        calls = []
        def fail(**kwargs):
            calls.append(kwargs)
            raise TimeoutError('secret provider body')
        self.ai.client.chat.completions.create = fail
        self.ai.config['max_retries'] = 1
        with patch('src.cv.time.sleep'), self.assertRaises(ValueError):
            self.service.analyze(self.sid)
        self.assertEqual(len(calls), 2)
        usage = rows(self.db, 'SELECT * FROM api_usage')
        self.assertEqual(len(usage), 2)
        self.assertTrue(all(r['estimated'] and r['output_tokens'] == 4500 for r in usage))
        self.assertNotIn('secret', encode(usage))
        self.ai.config['run_input_tokens'] = 1
        with patch('src.cv.time.sleep'), self.assertRaises(ValueError):
            self.service.analyze(self.sid)
        self.assertEqual(len(calls), 2)

    def test_docx_pdf_content_accents_and_pagination(self):
        _, _, vid = self.prepared()
        original = one(self.db, 'versions', vid)['data']
        original['sections'][0]['items'] *= 25
        long_vid = self.service.edit_version(vid, original)
        docx = Document(render_version(self.service, long_vid, 'docx'))
        text = '\n'.join(p.text for p in docx.paragraphs)
        self.assertIn('Élise Example', text)
        self.assertIn('Ingénieure logiciel', text)
        self.assertNotIn('Préciser les projets', text)
        pdf = PdfReader(render_version(self.service, long_vid, 'pdf'))
        self.assertGreater(len(pdf.pages), 1)
        pdf_text = '\n'.join(p.extract_text() for p in pdf.pages)
        self.assertIn('Élise Example', pdf_text)
        for item in original['sections'][0]['items'][:3]:
            self.assertIn(item['text'], pdf_text)
            self.assertIn(item['text'], text)

    def test_additive_legacy_migration(self):
        path = Path(self.temp.name) / 'legacy.db'
        conn = sqlite3.connect(path)
        conn.executescript('''CREATE TABLE jobs(id TEXT PRIMARY KEY,title TEXT,company TEXT,location TEXT,url TEXT,
            description TEXT,source TEXT,status TEXT,tier TEXT,ai_reason TEXT,date_added TEXT);
            INSERT INTO jobs VALUES('old','Title','Company','','','','fixture','APPROVED','1','keep','2020');
            CREATE TABLE api_usage(id INTEGER PRIMARY KEY,created_at TEXT,run_id TEXT,model TEXT,jobs INTEGER,
            input_tokens INTEGER,output_tokens INTEGER,cache_hit_tokens INTEGER,cache_miss_tokens INTEGER,
            estimated INTEGER,cost_usd REAL,latency REAL,error TEXT);''')
        conn.close()
        db = JobDatabase(path)
        self.assertEqual(rows(db, 'SELECT ai_reason FROM jobs')[0]['ai_reason'], 'keep')
        self.assertIn('operation', {r['name'] for r in rows(db, 'PRAGMA table_info(api_usage)')})
        db.close()
        db = JobDatabase(path)
        self.assertEqual(len(rows(db, 'SELECT * FROM jobs')), 1)
        db.close()

    def test_worker_restart_recovery_and_single_worker(self):
        with self.db.conn:
            self.db.conn.execute("INSERT INTO cv_tasks(operation,args,status) VALUES ('analyze',?,'RUNNING')", (encode({'source_id': self.sid}),))
        worker = CVWorker(self.service)
        worker.start()
        try:
            self.assertEqual(one(self.db, 'tasks', 1)['status'], 'INTERRUPTED')
            other = CVWorker(self.service)
            with self.assertRaises(RuntimeError):
                other.start()
            task_id = worker.retry(1)
            deadline = time.monotonic() + 5
            while one(self.db, 'tasks', task_id)['status'] in {'QUEUED', 'RUNNING'} and time.monotonic() < deadline:
                time.sleep(.02)
            self.assertEqual(one(self.db, 'tasks', task_id)['status'], 'COMPLETED')
        finally:
            worker.stop()


class WebTests(unittest.TestCase):
    def test_complete_api_workflow_and_csrf(self):
        from fastapi.testclient import TestClient
        from bs4 import BeautifulSoup
        from web import create_app
        provider = FakeProvider()
        with tempfile.TemporaryDirectory() as temp:
            app = create_app(Path(temp) / 'jobs.db', Path(temp) / 'private', {'max_retries': 0}, provider.client)
            with TestClient(app, base_url='http://127.0.0.1') as client:
                index = client.get('/')
                token = BeautifulSoup(index.text, 'html.parser').find('meta', attrs={'name': 'csrf-token'})['content']
                self.assertEqual(client.post('/api/source', json={}).status_code, 403)
                headers = {'x-csrf-token': token}
                self.assertEqual(client.post('/api/source', json={}, headers={**headers, 'Origin': 'https://evil.test'}).status_code, 403)
                self.assertEqual(client.get('/', headers={'Host': 'evil.test'}).status_code, 400)
                upload = client.post('/api/import', files={'file': ('cv.docx', docx_fixture())}, headers=headers)
                self.assertEqual(upload.status_code, 200)
                saved = client.post('/api/source', json=upload.json(), headers=headers).json()
                def task(operation, args):
                    body = {'operation': operation, 'args': args}
                    estimate = client.post('/api/estimate', json=body, headers=headers)
                    self.assertEqual(estimate.status_code, 200, estimate.text)
                    self.assertEqual(client.post('/api/tasks', json=body, headers=headers).status_code, 400)
                    body['estimate_key'] = estimate.json()['estimate_key']
                    tid = client.post('/api/tasks', json=body, headers=headers).json()['task_id']
                    deadline = time.monotonic() + 5
                    while time.monotonic() < deadline:
                        record = next(t for t in client.get('/api/tasks').json() if t['id'] == tid)
                        if record['status'] not in {'QUEUED', 'RUNNING'}:
                            self.assertEqual(record['status'], 'COMPLETED', record)
                            return json.loads(record['result'])
                        time.sleep(.02)
                    self.fail('Worker did not complete')
                pid = task('analyze', saved)['profile_id']
                c = categories(app.state.db)[0]
                cid = client.post(f'/api/categories/{pid}', json=[{**c['data'], 'id': c['id'], 'approved': True}], headers=headers).json()['category_ids'][0]
                invalid = client.post('/api/estimate', json={'operation': 'generate',
                    'args': {'category_id': cid, 'output_format': 'unknown'}}, headers=headers)
                self.assertEqual(invalid.status_code, 400)
                vid = task('generate', {'category_id': cid, 'output_format': 'latex'})['version_id']
                self.assertIn('Draft: not recommendable', client.get(f'/versions/{vid}').text)
                self.assertEqual(client.post(f'/api/versions/{vid}/approve', json={'reviewed': True}, headers=headers).status_code, 200)
                for extension in ['docx', 'pdf', 'tex']:
                    download = client.get(f'/versions/{vid}/download/{extension}')
                    self.assertEqual(download.status_code, 200, download.text[:200] if extension == 'pdf' and download.status_code != 200 else '')
                    self.assertIn('attachment', download.headers['content-disposition'])
                self.assertEqual(client.get(f'/versions/{vid}/download/exe').status_code, 400)


if __name__ == '__main__':
    unittest.main()
