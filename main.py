import os
import hashlib
import json
import time
import logging
import schedule
import concurrent.futures
import pandas as pd
from datetime import datetime
from dotenv import load_dotenv

from src.database import JobDatabase
from src.notifier import DiscordNotifier
from src.ai_filter import AIFilter
from src.scrapers.welcome_to_the_jungle import WelcomeToTheJungleScraper
from src.scrapers.linkedin import LinkedInScraper
from src.scrapers.hellowork import HelloWorkScraper
from src.scrapers.apec import ApecScraper
from src.scrapers.lesjeudis import LesJeudisScraper
from src.scrapers.francetravail import FranceTravailScraper

from src.configuration import load_config, build_criteria

logger = logging.getLogger(__name__)
USER_CONFIG = {}

db = None
notifier = None
ai = None


def initialize():
    global db, notifier, ai
    os.makedirs("logs", exist_ok=True)
    logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s',
                        handlers=[logging.FileHandler("logs/scraper.log", encoding="utf-8"), logging.StreamHandler()])
    load_dotenv()
    db = JobDatabase()
    notifier = DiscordNotifier()
    ai = AIFilter(USER_CONFIG.get("ai", {}), db=db)


def criteria_text():
    return build_criteria(USER_CONFIG or load_config())


def criteria_version():
    return hashlib.sha256(criteria_text().encode()).hexdigest()[:16]


def scrape_worker(scraper):
    """Worker function to run a single scraper in a thread."""
    logger.info(f"▶️ Starting scraper: {scraper.__class__.__name__}")
    try:
        jobs = scraper.scrape_jobs(db=db)
        counts = {"PENDING":0, "INVALID":0, "DUPLICATE":0}
        for job in jobs:
            status = db.add_job(job)
            if status in counts:
                counts[status] += 1
        with db.lock, db.conn:
            db.conn.execute("INSERT INTO scraper_runs(source,fetched,added,invalid,duplicates) VALUES (?,?,?,?,?)",
                (scraper.__class__.__name__,len(jobs),counts["PENDING"],counts["INVALID"],counts["DUPLICATE"]))
        logger.info("%s: %s", scraper.__class__.__name__, counts)
    except Exception as exc:
        logger.error("Scraper %s failed: %s", scraper.__class__.__name__, type(exc).__name__)
        with db.lock, db.conn:
            db.conn.execute("INSERT INTO scraper_runs(source,error) VALUES (?,?)", (scraper.__class__.__name__,type(exc).__name__))


def run_all_scrapers():
    classes = {"linkedin":LinkedInScraper,"hellowork":HelloWorkScraper,"apec":ApecScraper,
               "lesjeudis":LesJeudisScraper,"francetravail":FranceTravailScraper,
               "welcome_to_the_jungle":WelcomeToTheJungleScraper}
    profiles = USER_CONFIG["search_profiles"]
    sources = USER_CONFIG["sources"]
    # Each source runs its profiles sequentially; different sources run concurrently.
    def source_worker(source):
        for profile in profiles:
            scrape_worker(classes[source](keyword=profile["keyword"],location=profile.get("location",""),max_pages=profile.get("max_pages",10), **USER_CONFIG.get("source_options", {}).get(source, {})))
    with concurrent.futures.ThreadPoolExecutor(max_workers=min(5,len(sources)) or 1) as executor:
        list(executor.map(source_worker, sources))


def generate_run_summary(processed_records):
    """Generates an Excel spreadsheet report of all jobs processed in the current cycle."""
    if not processed_records:
        return None
        
    os.makedirs("reports", exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    filepath = f"reports/summary_{timestamp}.xlsx"
    
    # 1. Structure the data for the spreadsheet
    data = []
    for r in processed_records:
        job = r['job']
        data.append({
            "Status": r['status'],
            "Tier": str(r['tier']),
            "Title": job.title,
            "Company": job.company,
            "Location": job.location,
            "Source": job.source,
            "AI Reason": r['reason'],
            "URL": job.url
        })
        
    df = pd.DataFrame(data)
    
    # 2. Sort the data: APPROVED first, then sort by Tier
    df['Status_Sort'] = pd.Categorical(df['Status'], categories=["APPROVED", "REVIEW", "INVALID", "REJECTED", "ERROR"], ordered=True)
    df['Tier_Sort'] = pd.Categorical(df['Tier'], categories=["1", "2", "3", "4", "NONE"], ordered=True)
    df = df.sort_values(['Status_Sort', 'Tier_Sort']).drop(columns=['Status_Sort', 'Tier_Sort'])
    
    # 3. Export to Excel with styling
    import openpyxl
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Job Summary')
        worksheet = writer.sheets['Job Summary']
        
        # Freeze the top header row
        worksheet.freeze_panes = 'A2'
        
        # Define styles
        header_font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        header_fill = openpyxl.styles.PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        wrap_align = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
        
        # Format Headers
        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = openpyxl.styles.Alignment(horizontal='center', vertical='center')
            
        # Adjust column widths for aesthetics
        widths = {'A': 15, 'B': 10, 'C': 45, 'D': 25, 'E': 20, 'F': 15, 'G': 90, 'H': 30}
        for col, width in widths.items():
            worksheet.column_dimensions[col].width = width
            
        # Color coding and text wrapping
        tier_colors = {"1": "FFD700", "2": "C0C0C0", "3": "CD7F32", "4": "808080"}
        
        for row in range(2, worksheet.max_row + 1):
            for col in range(1, worksheet.max_column + 1):
                worksheet.cell(row=row, column=col).alignment = wrap_align
                
            # Color Status
            status_cell = worksheet.cell(row=row, column=1)
            if status_cell.value == "APPROVED":
                status_cell.fill = openpyxl.styles.PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
                status_cell.font = openpyxl.styles.Font(color="006100", bold=True)
            elif status_cell.value == "REJECTED":
                status_cell.fill = openpyxl.styles.PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
                status_cell.font = openpyxl.styles.Font(color="9C0006")
                
            # Color Tier
            tier_cell = worksheet.cell(row=row, column=2)
            tier_val = str(tier_cell.value)
            if tier_val in tier_colors:
                c = tier_colors[tier_val]
                tier_cell.fill = openpyxl.styles.PatternFill(start_color=c, end_color=c, fill_type="solid")
                tier_cell.font = openpyxl.styles.Font(bold=True)
                
            # Format URL as clickable
            url_cell = worksheet.cell(row=row, column=8)
            if url_cell.value and str(url_cell.value).startswith("http"):
                url_cell.hyperlink = url_cell.value
                url_cell.font = openpyxl.styles.Font(color="0563C1", underline="single")
            
    logger.info(f"📄 Generated verification spreadsheet: {filepath}")
    
    approved_count = len([r for r in processed_records if r['status'] == 'APPROVED'])
    return filepath, len(processed_records), approved_count

def retry_notifications():
    for row in db.pending_notifications():
        if notifier.send_job_alert(db.to_job(row), row['ai_reason'], row['tier']):
            db.notification_sent(row['id'])


def process_pending_jobs(dry_run=False):
    from src.content import invalid_reason
    ai.reset_run()
    criteria = criteria_text()
    pending = db.get_pending_jobs()
    eligible, records = [], []
    blacklist = USER_CONFIG.get('blacklist', {})
    for job in pending:
        reason = invalid_reason(job.description)
        status = 'INVALID' if reason else None
        if not status and any(c.casefold() in job.company.casefold() for c in blacklist.get('companies', []) if c):
            status, reason = 'REJECTED', 'Blacklisted company'
        # Keyword exclusions are explicit user policy, not inferred contract rules.
        if not status and any(k.casefold() in (job.title+' '+job.description).casefold() for k in blacklist.get('keywords', []) if k):
            status, reason = 'REJECTED', 'Blacklisted keyword'
        if status:
            if not dry_run:
                db.update_job_status(job.id,status,reason,criteria_version=criteria_version())
            records.append(dict(job=job,status=status,tier='NONE',reason=reason))
        else:
            eligible.append(job)
    batches, oversized = ai.plan(eligible, criteria)
    if dry_run:
        from src.ai_filter import estimate_tokens
        print(json.dumps({'pending':len(pending),'locally_filtered':len(records),'batches':len(batches),
            'oversized_review':len(oversized),'estimated_input_tokens':sum(estimate_tokens(ai.instructions(criteria))+estimate_tokens(ai.payload(b)) for b in batches),
            'reserved_output_tokens':sum(ai.output_budget(b) for b in batches),
            'run_input_limit':ai.config.get('run_input_tokens',500000),
            'run_output_limit':ai.config.get('run_output_tokens',50000),
            'estimator':'conservative UTF-8 byte estimate; actual API usage may be much lower'}, indent=2))
        return
    for job in oversized:
        reason = 'Description exceeds configured batch input budget; review or increase budget'
        db.update_job_status(job.id,'REVIEW',reason,criteria_version=criteria_version())
        records.append(dict(job=job,status='REVIEW',tier='NONE',reason=reason))
    for batch in batches:
        results = ai.evaluate_batch(batch, criteria)
        for job in batch:
            row = results.get(job.id)
            if row:
                reason = row['reason_code'] + ': ' + row['reason']
                if row.get('evidence'):
                    reason += ' | Evidence: ' + row['evidence']
                tier = str(row['tier']) if row['tier'] else 'NONE'
                db.update_job_status(job.id,row['decision'],reason,tier,criteria_version())
                records.append(dict(job=job,status=row['decision'],tier=tier,reason=reason))
            else:
                records.append(dict(job=job,status='ERROR',tier='NONE',reason='Budget exhausted or invalid API result; remains pending'))
    retry_notifications()
    result = generate_run_summary(records)
    if result:
        notifier.send_run_summary(*result)


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='JobFinder: configurable job discovery and ranking')
    parser.add_argument('--config', help='Use this YAML configuration instead of local/example defaults')
    parser.add_argument('--dry-run', action='store_true', help='Plan pending jobs without API calls or notifications')
    parser.add_argument('--once', action='store_true')
    parser.add_argument('--process-only', action='store_true')
    args = parser.parse_args()
    USER_CONFIG = load_config(args.config)
    initialize()
    if args.dry_run:
        process_pending_jobs(dry_run=True)
    else:
        if not args.process_only:
            run_all_scrapers()
        process_pending_jobs()
        if not args.once:
            schedule.every(6).hours.do(run_all_scrapers)
            schedule.every(15).minutes.do(process_pending_jobs)
            while True:
                schedule.run_pending()
                time.sleep(30)
